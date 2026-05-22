"""Generate Lightworks Pro capability and command reference Word document."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

BRAND_BLUE  = RGBColor(0x00, 0x55, 0xB8)
BRAND_DARK  = RGBColor(0x1A, 0x1A, 0x2E)
LIGHT_GREY  = RGBColor(0xF4, 0xF6, 0xF9)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
RED_LIMIT   = RGBColor(0xC0, 0x39, 0x2B)
GREEN_CAP   = RGBColor(0x1E, 0x88, 0x55)

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin   = Inches(1.0)
    section.right_margin  = Inches(1.0)

# ── Helpers ───────────────────────────────────────────────────────────────────

def heading1(text):
    p = doc.add_heading(text, level=1)
    p.runs[0].font.color.rgb = BRAND_BLUE
    p.runs[0].font.size = Pt(16)
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after  = Pt(6)
    return p

def heading2(text):
    p = doc.add_heading(text, level=2)
    p.runs[0].font.color.rgb = BRAND_DARK
    p.runs[0].font.size = Pt(13)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)
    return p

def body(text, bold=False, colour=None):
    p = doc.add_paragraph(text)
    r = p.runs[0]
    r.font.size = Pt(11)
    r.font.bold = bold
    if colour:
        r.font.color.rgb = colour
    p.paragraph_format.space_after = Pt(3)
    return p

def bullet(text, level=0, colour=None):
    p = doc.add_paragraph(text, style="List Bullet")
    p.runs[0].font.size = Pt(11)
    if colour:
        p.runs[0].font.color.rgb = colour
    p.paragraph_format.left_indent = Inches(0.3 + level * 0.2)
    p.paragraph_format.space_after = Pt(2)
    return p

def cmd_row(table, command, description):
    row = table.add_row()
    c = row.cells[0].paragraphs[0]
    c.clear()
    r = c.add_run(command)
    r.font.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = BRAND_BLUE
    d = row.cells[1].paragraphs[0]
    d.clear()
    rd = d.add_run(description)
    rd.font.size = Pt(10.5)

def section_table(headers, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, (h, w) in enumerate(zip(headers, widths)):
        hdr[i].width = Inches(w)
        p = hdr[i].paragraphs[0]
        p.clear()
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = WHITE
        tc = hdr[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "0055B8")
        tcPr.append(shd)
    return t

def cmd_table(headers=("Voice Command / Hotkey", "What It Does"), widths=(2.8, 3.8)):
    return section_table(headers, widths)

def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Inches(0.3)
    p.paragraph_format.space_after  = Pt(6)
    r = p.add_run("ℹ  " + text)
    r.font.size   = Pt(10)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x5A, 0x5A, 0x7A)
    return p

def hr():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pb  = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "6")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), "CCCCCC")
    pb.append(bot)
    pPr.append(pb)

# ══════════════════════════════════════════════════════════════════════════════
# TITLE PAGE
# ══════════════════════════════════════════════════════════════════════════════

title = doc.add_heading("Lightworks Pro", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.color.rgb = BRAND_BLUE
title.runs[0].font.size = Pt(28)

sub = doc.add_paragraph("Capability Guide & Voice Command Reference")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.size = Pt(14)
sub.runs[0].font.color.rgb = BRAND_DARK

org = doc.add_paragraph("Abu Hoque  ·  Lighthouse Guild, New York")
org.alignment = WD_ALIGN_PARAGRAPH.CENTER
org.runs[0].font.size = Pt(11)
org.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

dt = doc.add_paragraph(datetime.date.today().strftime("Version 1.2.0  ·  %B %Y"))
dt.alignment = WD_ALIGN_PARAGRAPH.CENTER
dt.runs[0].font.size = Pt(10)
dt.runs[0].font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

doc.add_paragraph()
hr()
doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 1. OVERVIEW
# ══════════════════════════════════════════════════════════════════════════════

heading1("1.  Overview")
body(
    "Lightworks Pro is a voice-controlled productivity assistant designed "
    "for blind and low-vision professionals. It runs silently in the Windows "
    "system tray and responds to voice commands and global hotkeys — no mouse "
    "or screen navigation required. Integration with Microsoft 365 unlocks "
    "email, calendar, Teams calling, and more."
)
note(
    "Requires: Windows 10/11 (64-bit), Microsoft Teams (new Teams 2024), "
    "Microsoft 365 account with Teams Phone licence for PSTN/extension calling."
)

# ══════════════════════════════════════════════════════════════════════════════
# 2. OUTLOOK / EMAIL
# ══════════════════════════════════════════════════════════════════════════════

heading1("2.  Outlook / Email")

heading2("Capabilities")
for cap in [
    "Reads up to 10 unread inbox messages by voice — sender, subject, and preview.",
    "Full message read-back: say 'next' to move to the next message, 'read full' for the complete body.",
    "Compose and send a new email entirely by voice (dictate To, Subject, and body).",
    "Reply to the current message by voice — Lightworks dictates your reply and sends it.",
    "Real-time inbox monitoring — announces new high-priority emails as they arrive.",
    "Briefing mode ('what did I miss') — combines recent emails and upcoming meetings in one report.",
]:
    bullet(cap, colour=GREEN_CAP)

heading2("Limitations")
for lim in [
    "Reads the inbox only — Sent Items, Drafts, and sub-folders are not accessible by voice.",
    "Attachment reading is not supported — attachments are announced by name only.",
    "Email search ('find email from John') is not yet implemented.",
    "Inline images and HTML-only emails are read as plain text; formatting is stripped.",
    "Composing supports one To: address per message (no CC/BCC by voice).",
    "Shared mailboxes are not supported — only the signed-in user's primary mailbox.",
]:
    bullet(lim, colour=RED_LIMIT)

heading2("Voice Commands — Email")
t = cmd_table()
cmds = [
    ("Ctrl+Shift+F2  /  \"inbox\"",      "Read up to 10 unread emails"),
    ("\"next\"",                          "Move to the next email while in inbox"),
    ("\"read full\"",                     "Read the complete body of the current email"),
    ("\"reply\"",                         "Start a voice reply to the current email"),
    ("\"stop\"",                          "Stop reading and return to standby"),
    ("\"compose email to [name] about [subject]\"",
                                          "Open a new email — Lightworks will ask you to dictate the body"),
    ("\"send email to [name] about [subject]\"",
                                          "Same as compose email"),
    ("\"teams message\" / \"teams chat\"","Read recent unread Teams chat messages"),
    ("\"what did I miss\"",               "Briefing: recent emails + upcoming meetings"),
]
for c, d in cmds:
    cmd_row(t, c, d)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 3. CALENDAR
# ══════════════════════════════════════════════════════════════════════════════

heading1("3.  Calendar")

heading2("Capabilities")
for cap in [
    "Reads your Outlook and Teams calendar — both share the same Microsoft 365 calendar.",
    "On-demand: tells you your next upcoming meeting with time, platform, subject, and your role (host or participant).",
    "Daily schedule: reads all meetings for today, tomorrow, or any named weekday.",
    "Automatic meeting alerts at 10 minutes and 5 minutes before each meeting.",
    "Voice consent to join: asks 'Should I join?' — say yes to open Teams or Zoom automatically.",
    "Detects Teams, Zoom, Webex, and Google Meet join links from meeting invites.",
    "Sets your Teams presence to Busy / In a Meeting on join; resets to Available after.",
    "Briefing mode ('what did I miss') includes your upcoming meetings.",
]:
    bullet(cap, colour=GREEN_CAP)

heading2("Limitations")
for lim in [
    "Read-only — cannot create, edit, or delete calendar events by voice.",
    "Reads the primary calendar only — shared calendars and room calendars are not shown.",
    "Meeting alerts look ahead 24 hours — events further than 24 hours away are not pre-announced.",
    "Recurring event exceptions (e.g. a moved instance) may show the original time.",
    "Google Meet and Webex meetings open in the default browser (no native app deep link on Windows).",
]:
    bullet(lim, colour=RED_LIMIT)

heading2("Voice Commands — Calendar")
t = cmd_table()
cmds = [
    ("Ctrl+Shift+F3  /  \"next meeting\"",  "Read your next upcoming meeting"),
    ("\"calendar\"  /  \"my meetings\"",     "Same as next meeting"),
    ("\"today\"",                            "Read all of today's meetings"),
    ("\"tomorrow\"",                         "Read all of tomorrow's meetings"),
    ("\"Monday\" … \"Sunday\"",              "Read meetings for that day of the week"),
    ("\"upcoming\"  /  \"schedule\"",        "Read upcoming meetings"),
    ("\"yes\"  (when alert fires)",          "Confirm joining the meeting Teams/Zoom opens automatically"),
    ("\"no\"  (when alert fires)",           "Dismiss the meeting alert"),
    ("\"remind me in [N] minutes about [topic]\"",
                                             "Set a voice reminder (separate from calendar)"),
]
for c, d in cmds:
    cmd_row(t, c, d)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 4. TEAMS MEETINGS (JOINING & IN-MEETING CONTROLS)
# ══════════════════════════════════════════════════════════════════════════════

heading1("4.  Teams Meetings — Joining & In-Meeting Controls")

heading2("Joining a Meeting")
for cap in [
    "Automatic alert at 10 and 5 minutes before any Teams meeting in your calendar.",
    "Say 'yes' to join — Lightworks opens Teams via the native deep link (msteams://l/meetup-join/…).",
    "Auto-mute on join: microphone is muted immediately after Teams connects (if configured).",
    "Teams presence is set to Busy / In a Meeting automatically on join.",
    "Presence resets to Available 5 minutes after the scheduled meeting end time.",
    "Works for meetings you organised (host) and meetings you were invited to (participant).",
]:
    bullet(cap, colour=GREEN_CAP)

heading2("In-Meeting Voice Controls")
for cap in [
    "Mute / unmute microphone by voice.",
    "Toggle camera on or off.",
    "Raise or lower hand.",
    "Leave the meeting — Teams hang-up shortcut is sent automatically.",
]:
    bullet(cap, colour=GREEN_CAP)

heading2("Limitations — Meetings")
for lim in [
    "Cannot create or schedule a new Teams meeting by voice.",
    "Cannot share screen, manage participants, or use the chat panel by voice.",
    "In-meeting controls require the Teams window to be visible (not minimised to taskbar).",
    "Auto-join only triggers for meetings in your primary Microsoft 365 calendar.",
    "Webex and Google Meet meetings open in the browser — no in-meeting voice controls for those platforms.",
]:
    bullet(lim, colour=RED_LIMIT)

heading2("Voice Commands — Teams Meetings")
t = cmd_table()
cmds = [
    ("\"mute me\"  /  \"mute\"  /  \"unmute\"",  "Toggle microphone on/off in Teams or Zoom"),
    ("\"camera\"  /  \"video on\"  /  \"video off\"",
                                                    "Toggle camera on/off"),
    ("\"raise hand\"  /  \"hand up\"",             "Raise your hand in the meeting"),
    ("\"leave meeting\"  /  \"hang up\"  /  \"end meeting\"",
                                                    "Leave the current meeting"),
    ("\"yes\"  (when alert fires)",                 "Join the meeting — Teams opens automatically"),
    ("\"no\"  /  \"stop\"  (when alert fires)",     "Dismiss the meeting alert"),
]
for c, d in cmds:
    cmd_row(t, c, d)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 5. TEAMS CALLING
# ══════════════════════════════════════════════════════════════════════════════

heading1("5.  Teams Calling")

heading2("How It Works")
body(
    "Lightworks Pro initiates calls using registered Windows URI protocols "
    "('callto:', 'tel:') that Teams handles natively — no keyboard automation "
    "is used to navigate Teams. When Teams shows its 'Would you like to call?' "
    "confirmation dialog, Lightworks automatically confirms it using Windows UI "
    "Automation (UIA Invoke), which is fully compatible with JAWS, NVDA, Fusion, "
    "and ZoomText because it activates the button through the accessibility layer "
    "with no keyboard events generated."
)

heading2("Capabilities")
for cap in [
    "Call a person by name — Lightworks looks them up in your Microsoft 365 contacts/directory.",
    "Video call a person by name.",
    "Dial any full phone number (PSTN) — requires Teams Phone licence.",
    "Dial an internal office extension — requires Teams Phone direct routing.",
    "Conference call with multiple people (say 'conference call with John and Mary').",
    "Answer an incoming Teams call by voice.",
    "Decline an incoming Teams call by voice.",
    "Contact search covers three tiers: personal Outlook contacts → org People directory → Azure AD user search.",
    "AT-safe confirmation dialog: auto-clicks 'Call' via UIA Invoke — no conflict with screen readers.",
]:
    bullet(cap, colour=GREEN_CAP)

heading2("Limitations")
for lim in [
    "PSTN calls (phone numbers) require a Teams Phone calling plan or direct routing licence — IT must configure this.",
    "Extensions require Teams Phone with extension routing configured in the tenant.",
    "Video calls via msteams:// deep link: if Teams rejects the URI, it falls back to audio-only (callto:).",
    "Conference calls use the msteams:// deep link — if Teams' internal SkypeMigration handler rejects it, conference may not connect.",
    "Contact search requires Microsoft 365 sign-in and at least one of: People.Read consent, Contacts.Read, or User.ReadBasic.All scope.",
    "Calling non-Teams users (people without Teams) requires their full PSTN number — name lookup only returns Teams/M365 accounts.",
    "Zoom calls are not initiated by voice — only Teams calling is supported for outgoing calls.",
]:
    bullet(lim, colour=RED_LIMIT)

heading2("Voice Commands — Teams Calling")
t = cmd_table()
cmds = [
    ("Ctrl+Shift+F4  /  \"teams call\"",   "Open the Teams call prompt"),
    ("\"call [name]\"",                    "Audio call to a contact by name via Teams"),
    ("\"video call [name]\"  /  \"video chat [name]\"",
                                           "Video call to a contact by name"),
    ("\"dial [phone number]\"",            "PSTN call to a full phone number (e.g. dial +1 929 245 9540)"),
    ("\"extension [number]\"  /  \"ext [number]\"",
                                           "Dial an internal office extension"),
    ("\"conference call with [name] and [name]\"",
                                           "Start a Teams conference call with multiple people"),
    ("\"answer\"  /  \"answer call\"",     "Answer an incoming Teams call"),
    ("\"answer with video\"",              "Answer an incoming Teams call with video on"),
    ("\"decline\"  /  \"decline call\"",   "Decline an incoming Teams call"),
    ("\"hang up\"  /  \"leave call\"",     "End the active call"),
]
for c, d in cmds:
    cmd_row(t, c, d)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 6. COMPLETE VOICE COMMAND REFERENCE
# ══════════════════════════════════════════════════════════════════════════════

heading1("6.  Complete Voice Command Reference")
note(
    "Activate voice input: press Ctrl+Shift+F1 (or use the tray icon), "
    "wait for the beep, then speak clearly. You have 15 seconds. "
    "Press Ctrl+Shift+F1 again at any time to stop speech immediately."
)

heading2("Global Hotkeys")
t = cmd_table()
cmds = [
    ("Ctrl + Shift + F1",  "Open voice input — speak a command after the beep"),
    ("Ctrl + Shift + F2",  "Read inbox (unread emails)"),
    ("Ctrl + Shift + F3",  "Read next meeting"),
    ("Ctrl + Shift + F4",  "Teams call — prompts for a contact name or number"),
]
for c, d in cmds:
    cmd_row(t, c, d)

doc.add_paragraph()
heading2("System & Utilities")
t = cmd_table()
cmds = [
    ("\"what time is it\"  /  \"time\"",      "Speaks the current time"),
    ("\"what's today's date\"  /  \"date\"",   "Speaks today's date"),
    ("\"battery\"",                            "Speaks battery level and charging status"),
    ("\"where am I\"  /  \"active window\"",   "Announces the currently focused application"),
    ("\"recent files\"",                       "Lists recently opened documents"),
    ("\"read clipboard\"  /  \"clipboard\"",   "Reads whatever text is on the clipboard"),
    ("\"spell [word]\"",                       "Spells a word using the NATO phonetic alphabet"),
    ("\"weather\"  /  \"temperature\"",        "Reads the current local weather"),
    ("\"faster\"  /  \"slower\"  /  \"normal speed\"",
                                               "Adjusts the text-to-speech speaking rate"),
    ("\"repeat that\"  /  \"say again\"",      "Repeats the last spoken announcement"),
    ("\"help\"  /  \"commands\"",              "Reads all available commands"),
    ("\"test\"",                               "Plays a test announcement to verify audio"),
    ("\"quit\"  /  \"exit\"  /  \"goodbye\"",  "Closes Lightworks Pro"),
]
for c, d in cmds:
    cmd_row(t, c, d)

doc.add_paragraph()
heading2("Email & Messaging")
t = cmd_table()
cmds = [
    ("\"inbox\"  /  \"check email\"  /  \"my email\"",  "Read up to 10 unread emails"),
    ("\"next\"",                                         "Next email (while in inbox)"),
    ("\"read full\"",                                    "Read complete body of current email"),
    ("\"reply\"",                                        "Voice reply to current email"),
    ("\"stop\"",                                         "Stop reading"),
    ("\"compose email to [name] about [subject]\"",      "Compose and send a new email by voice"),
    ("\"send email to [name] about [subject]\"",         "Same as compose email"),
    ("\"teams message\"  /  \"teams chat\"",             "Read recent unread Teams chat messages"),
    ("\"what did I miss\"  /  \"briefing\"",             "Recent emails + upcoming meetings summary"),
]
for c, d in cmds:
    cmd_row(t, c, d)

doc.add_paragraph()
heading2("Calendar")
t = cmd_table()
cmds = [
    ("\"next meeting\"  /  \"my meetings\"  /  \"calendar\"",  "Your next upcoming meeting"),
    ("\"today\"",                                               "All meetings today"),
    ("\"tomorrow\"",                                            "All meetings tomorrow"),
    ("\"Monday\" … \"Sunday\"",                                 "Meetings on that day of the week"),
    ("\"upcoming\"  /  \"schedule\"",                           "Upcoming meetings in the next 24 hours"),
    ("\"yes\"  (on meeting alert)",                             "Join the meeting automatically"),
    ("\"no\"  (on meeting alert)",                              "Dismiss the alert"),
]
for c, d in cmds:
    cmd_row(t, c, d)

doc.add_paragraph()
heading2("Calls & Meeting Controls")
t = cmd_table()
cmds = [
    ("\"call [name]\"",                                     "Audio call via Teams"),
    ("\"video call [name]\"",                               "Video call via Teams"),
    ("\"dial [number]\"  /  \"phone number [number]\"",     "PSTN call via Teams Phone"),
    ("\"extension [number]\"  /  \"ext [number]\"",         "Internal extension via Teams Phone"),
    ("\"conference call with [name] and [name]\"",          "Teams conference call"),
    ("\"answer\"  /  \"answer call\"",                      "Answer incoming Teams call"),
    ("\"answer with video\"",                               "Answer with video"),
    ("\"decline\"  /  \"decline call\"",                    "Decline incoming call"),
    ("\"mute me\"  /  \"mute\"  /  \"unmute\"",             "Toggle microphone (Teams or Zoom)"),
    ("\"camera\"  /  \"video on\"  /  \"video off\"",       "Toggle camera"),
    ("\"raise hand\"",                                      "Raise hand in meeting"),
    ("\"leave meeting\"  /  \"hang up\"  /  \"end meeting\"",  "Leave current meeting/call"),
]
for c, d in cmds:
    cmd_row(t, c, d)

doc.add_paragraph()
heading2("Tasks, Notes & Contacts")
t = cmd_table()
cmds = [
    ("\"my tasks\"  /  \"to do list\"",                         "Read Microsoft To Do task list"),
    ("\"add task [title]\"  /  \"new task [title]\"",           "Add a new task to Microsoft To Do"),
    ("\"complete task [name]\"  /  \"mark done [name]\"",       "Mark a task as complete"),
    ("\"create note\"  /  \"dictate note\"  /  \"take note\"",  "Save a voice note to OneNote"),
    ("\"who is [name]\"  /  \"find contact [name]\"",           "Look up a contact's details"),
]
for c, d in cmds:
    cmd_row(t, c, d)

doc.add_paragraph()
heading2("Status & Notifications")
t = cmd_table()
cmds = [
    ("\"set status busy\"  /  \"I'm busy\"",        "Set Teams presence to Busy"),
    ("\"set status available\"  /  \"I'm available\"",  "Set Teams presence to Available"),
    ("\"do not disturb\"",                          "Set Teams presence to Do Not Disturb"),
    ("\"quiet for [N] minutes\"",                   "Pause all voice announcements for N minutes"),
    ("\"terse mode\"  /  \"brief mode\"",           "Shorter, faster announcements"),
    ("\"verbose mode\"",                            "Full detailed announcements"),
]
for c, d in cmds:
    cmd_row(t, c, d)

doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# 7. SETUP REQUIREMENTS
# ══════════════════════════════════════════════════════════════════════════════

heading1("7.  One-Time Setup Requirements")
body("The following steps are required once after installation:")

steps = [
    ("Install Lightworks Pro",
     "Run LightworksPro_Setup_1.2.0.exe. The app starts automatically in the system tray."),
    ("Register Microsoft 365",
     "In the Start Menu → Lightworks Pro → 'Register Microsoft 365'. "
     "A browser window opens — sign in with your work account. "
     "OR run: scripts\\register_azure_app.ps1 in PowerShell 7."),
    ("Grant admin consent (if required)",
     "If contact search returns empty results, your IT admin must grant "
     "admin consent for People.Read. Run: scripts\\grant_admin_consent.py "
     "and share the URL with your admin. User.ReadBasic.All usually "
     "auto-consents and provides org directory search as a fallback."),
    ("Teams Phone (for PSTN/extensions)",
     "Your Microsoft 365 account must have a Teams Phone calling plan or "
     "direct routing licence assigned by IT. "
     "Without this, audio/video calls to names still work; "
     "phone numbers and extensions do not."),
]
for i, (title, desc) in enumerate(steps, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r1 = p.add_run(f"Step {i}: {title} — ")
    r1.font.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run(desc)
    r2.font.size = Pt(11)

hr()
foot = doc.add_paragraph(
    f"Lightworks Pro v1.2.0  ·  Abu Hoque & Lighthouse Guild, New York  ·  MIT License  ·  "
    f"{datetime.date.today().strftime('%B %Y')}"
)
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
foot.runs[0].font.size  = Pt(9)
foot.runs[0].font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)

# ── Save ──────────────────────────────────────────────────────────────────────
out = r"c:\Claude Code\Lightworks Pro\dist\LightworksPro_Guide.docx"
doc.save(out)
print(f"Saved: {out}")
